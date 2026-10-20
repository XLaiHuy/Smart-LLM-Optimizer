"""Bounded parser subprocess; stdout is private IPC, never application logging."""

import json
import sys
import zipfile
from pathlib import Path

from docx import Document as WordDocument
from pypdf import PdfReader


def extract(path: Path, suffix: str, max_pages: int, max_chars: int) -> list[tuple[str, str]]:
    sections: list[tuple[str, str]] = []
    if suffix in (".txt", ".md"):
        sections = [("text", path.read_text(encoding="utf-8-sig"))]
    elif suffix == ".pdf":
        if not path.read_bytes()[:1024].lstrip().startswith(b"%PDF-"):
            raise ValueError("FILE_TYPE_INVALID")
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("ENCRYPTED_DOCUMENT")
        if len(reader.pages) > max_pages:
            raise ValueError("DOCUMENT_TOO_LARGE")
        sections = [
            (f"page {i + 1}", page.extract_text() or "") for i, page in enumerate(reader.pages)
        ]
    elif suffix == ".docx":
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if (
                len(entries) > 2000
                or sum(i.file_size for i in entries) > 20 * 1024 * 1024
                or "word/document.xml" not in archive.namelist()
            ):
                raise ValueError("DOCUMENT_TOO_LARGE")
            if any(i.file_size > max(1, i.compress_size) * 200 for i in entries):
                raise ValueError("DOCUMENT_TOO_LARGE")
        doc = WordDocument(str(path))
        sections = [
            (f"paragraph {i + 1}", paragraph.text) for i, paragraph in enumerate(doc.paragraphs)
        ]
        for i, table in enumerate(doc.tables):
            sections.append(
                (
                    f"table {i + 1}",
                    "\n".join(" | ".join(c.text for c in r.cells) for r in table.rows),
                )
            )
    else:
        raise ValueError("UNSUPPORTED_FILE_TYPE")
    if sum(len(text) for _, text in sections) > max_chars:
        raise ValueError("DOCUMENT_TOO_LARGE")
    if not any(text.strip() for _, text in sections):
        raise ValueError("NO_EXTRACTABLE_TEXT")
    if any("\x00" in text for _, text in sections):
        raise ValueError("FILE_TYPE_INVALID")
    return sections


if __name__ == "__main__":
    try:
        print(
            json.dumps(
                {
                    "sections": extract(
                        Path(sys.argv[1]), sys.argv[2], int(sys.argv[3]), int(sys.argv[4])
                    )
                },
                ensure_ascii=True,
            )
        )
    except Exception as exc:
        allowed = {
            "DOCUMENT_TOO_LARGE",
            "FILE_TYPE_INVALID",
            "ENCRYPTED_DOCUMENT",
            "NO_EXTRACTABLE_TEXT",
            "UNSUPPORTED_FILE_TYPE",
        }
        print(json.dumps({"error": str(exc) if str(exc) in allowed else "DOCUMENT_PARSE_FAILED"}))
        sys.exit(1)
